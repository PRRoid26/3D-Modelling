"""
Borehole Log - Automated Colour & Texture Assignment
=====================================================
Reads the Stratigraphy column of a borehole log Excel file and
automatically fills both the Colour and Texture columns:

  Colour  – solid background fill derived from:
              1. Colour descriptor phrases in the stratigraphy text
              2. Lithology keyword fallback (CLAY, SAND, SILT, ...)

  Texture – Excel pattern fill (hatching) derived from the Index.xlsx legend:
              Clay        -> lightGray
              Sand        -> darkGray
              Silt        -> gray0625
              Calcarenite -> lightUp
              Sandstone   -> lightTrellis

Usage:
    python automate_borehole_colours.py <input.xlsx> [output.xlsx]

    If output path is omitted, the file is saved as <input>_coloured.xlsx

Dependencies:
    pip install pandas openpyxl
"""

import sys
import shutil
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import PatternFill


# =============================================================================
# COLOUR MAPPINGS  (edit to add / change colours)
# =============================================================================

# Colour descriptor phrases found in stratigraphy text -> solid hex fill colour.
# Longer phrases are matched first to prevent partial-match conflicts.
COLOUR_PHRASE_MAP = {
    "dark greenish grey":           "7A9E9F",
    "light greenish grey":          "B0C8B8",
    "dark greenish":                "4E7060",
    "greenish grey":                "9AB0A0",
    "light brownish grey":          "C3B9A8",
    "light olive brown":            "B5A55A",
    "olive yellow":                 "C8C056",
    "olive brown":                  "8B7D3A",
    "yellowish brown":              "C9A84C",
    "pale yellow":                  "F5E6A0",
    "light grey":                   "D0D8D8",
    "dark grey":                    "6E7B7B",
    "grey":                         "AABABA",
}

# Capitalised lithology keyword in stratigraphy -> solid hex fill colour.
# Used only when no colour phrase is matched above.
LITHOLOGY_COLOUR_MAP = {
    "CLAY":        "A8C5C5",   # muted teal-blue
    "SAND":        "F0D890",   # sandy yellow
    "SILT":        "C8B478",   # olive-tan
    "CALCARENITE": "E8D8B0",   # cream
    "SANDSTONE":   "D4975A",   # orange-brown
    "GRAVEL":      "D6C89A",   # warm beige
    "LIMESTONE":   "E8E0C8",   # pale warm grey
    "SHALE":       "9E9E7E",   # olive grey
}

DEFAULT_COLOUR = "FFFFFF"   # white - fallback when nothing matches


# =============================================================================
# TEXTURE MAPPINGS  (as defined in Index.xlsx column B pattern fills)
# =============================================================================
# Maps lithology keyword -> (Excel fill_type, fgColor hex, bgColor hex)
#   fill_type = Excel hatch pattern name
#   fgColor   = hatch / dot colour
#   bgColor   = background colour behind the hatching

LITHOLOGY_TEXTURE_MAP = {
    #  keyword        fill_type         fgColor    bgColor
    "CLAY":        ("lightGray",    "000000",  "A8C5C5"),
    "SAND":        ("darkGray",     "000000",  "F0D890"),
    "SILT":        ("gray0625",     "000000",  "C8B478"),
    "CALCARENITE": ("lightUp",      "000000",  "E8D8B0"),
    "SANDSTONE":   ("lightTrellis", "000000",  "D4975A"),
    "GRAVEL":      ("darkTrellis",  "000000",  "D6C89A"),
    "LIMESTONE":   ("lightGrid",    "000000",  "E8E0C8"),
    "SHALE":       ("lightDown",    "000000",  "9E9E7E"),
}


# =============================================================================
# CORE LOGIC
# =============================================================================

_SORTED_PHRASES   = sorted(COLOUR_PHRASE_MAP.items(),    key=lambda x: -len(x[0]))
_SORTED_LITHO_COL = sorted(LITHOLOGY_COLOUR_MAP.items(), key=lambda x: -len(x[0]))
_SORTED_LITHO_TEX = sorted(LITHOLOGY_TEXTURE_MAP.items(), key=lambda x: -len(x[0]))


def assign_colour(stratigraphy):
    """
    Return a 6-char hex colour string for a stratigraphy description.

    Logic:
        1. Scan for colour descriptor phrases (longest match wins).
        2. Detect capitalised lithology keyword as fallback (longest match wins,
           so SANDSTONE is checked before SAND).
        3. Return DEFAULT_COLOUR if nothing matches.
        4. Return None for empty / NaN values.
    """
    if pd.isna(stratigraphy) or str(stratigraphy).strip() == "":
        return None

    text = str(stratigraphy)

    for phrase, hex_col in _SORTED_PHRASES:
        if phrase.lower() in text.lower():
            return hex_col

    for keyword, hex_col in _SORTED_LITHO_COL:
        if keyword in text:
            return hex_col

    return DEFAULT_COLOUR


def assign_texture(stratigraphy):
    """
    Return (fill_type, fgColor, bgColor) for the Texture cell, or None if
    the stratigraphy is empty / no lithology keyword is matched.

    Keywords are checked longest-first so SANDSTONE is never shadowed by SAND.
    """
    if pd.isna(stratigraphy) or str(stratigraphy).strip() == "":
        return None

    text = str(stratigraphy)

    for keyword, texture in _SORTED_LITHO_TEX:
        if keyword in text:
            return texture

    return None


# =============================================================================
# EXCEL PROCESSING
# =============================================================================

def _find_headers(ws, *column_names):
    """Scan the workbook for header cells. Returns (header_row, {name: col})."""
    result = {}
    header_row = None
    for row in ws.iter_rows():
        for cell in row:
            if cell.value in column_names:
                result[cell.value] = cell.column
                header_row = cell.row
        if len(result) == len(column_names):
            break
    return header_row, result


def process_file(input_path, output_path=None):
    """
    Read input Excel file, assign colours and textures from the Stratigraphy
    column, write results into the Colour and Texture columns, and save.

    Returns the path of the saved file.
    """
    input_path = Path(input_path)
    if output_path is None:
        output_path = input_path.with_stem(input_path.stem + "_coloured")
    output_path = Path(output_path)

    shutil.copy(input_path, output_path)

    df = pd.read_excel(input_path)

    required = {"Stratigraphy", "Colour", "Texture"}
    missing = required - set(df.columns)
    if missing:
        raise ValueError(f"Missing expected column(s): {missing}")

    colours  = df["Stratigraphy"].apply(assign_colour)
    textures = df["Stratigraphy"].apply(assign_texture)

    wb = load_workbook(output_path)
    ws = wb.active

    header_row, cols = _find_headers(ws, "Stratigraphy", "Colour", "Texture")

    if header_row is None:
        raise ValueError("Could not locate column headers in the workbook.")

    missing_headers = required - set(cols.keys())
    if missing_headers:
        raise ValueError(f"Could not find header(s) in workbook: {missing_headers}")

    colour_col  = cols["Colour"]
    texture_col = cols["Texture"]

    for i in range(len(df)):
        excel_row = header_row + 1 + i

        # -- Colour cell -------------------------------------------------------
        hex_col = colours.iloc[i]
        colour_cell = ws.cell(row=excel_row, column=colour_col)
        if hex_col:
            colour_cell.value = f"#{hex_col}"
            colour_cell.fill  = PatternFill("solid", fgColor=hex_col)
        else:
            colour_cell.value = None
            colour_cell.fill  = PatternFill()

        # -- Texture cell ------------------------------------------------------
        texture = textures.iloc[i]
        tex_cell = ws.cell(row=excel_row, column=texture_col)
        if texture:
            fill_type, fg, bg = texture
            tex_cell.fill = PatternFill(
                fill_type=fill_type,
                fgColor=fg,
                bgColor=bg,
            )
        else:
            tex_cell.fill = PatternFill()

    wb.save(output_path)
    return str(output_path)


# =============================================================================
# ENTRY POINT
# =============================================================================

if __name__ == "__main__":
    if len(sys.argv) < 2:
        print(__doc__)
        sys.exit(1)

    input_file  = sys.argv[1]
    output_file = sys.argv[2] if len(sys.argv) > 2 else None

    try:
        saved = process_file(input_file, output_file)
        print(f"Done. File saved to: {saved}")
    except Exception as e:
        print(f"Error: {e}")
        sys.exit(1)
